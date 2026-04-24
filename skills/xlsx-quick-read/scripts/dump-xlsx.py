#!/usr/bin/env python
"""
Dump xlsx cell data in CSV format.
AI agent utility for quick xlsx inspection.

Usage:
    python dump-xlsx.py <file>
    python dump-xlsx.py <file> --sheet=<name>
    python dump-xlsx.py <file> --sheet=<name> --cols=A:C --rows=1:10
    python dump-xlsx.py <file> --sheet=<name> --cols=A,F,X:Z,BA:BD --rows=1:3,5:9
"""

import csv
import sys
from datetime import datetime

from openpyxl import load_workbook  # type: ignore
from openpyxl.utils import column_index_from_string, get_column_letter  # type: ignore


def parse_col_ranges(spec: str) -> set[int]:
    """Parse comma-separated column ranges (e.g. 'A,F,X:Z,BA:BD') into 1-based column indices."""
    indices = set()
    for part in spec.split(","):
        part = part.strip().upper()
        if ":" in part:
            start, end = part.split(":", 1)
            start_idx = column_index_from_string(start)
            end_idx = column_index_from_string(end)
            indices.update(range(start_idx, end_idx + 1))
        else:
            indices.add(column_index_from_string(part))
    return indices


def parse_row_ranges(spec: str) -> set[int]:
    """Parse comma-separated row ranges (e.g. '1:3,5:9') into 1-based row numbers."""
    rows = set()
    for part in spec.split(","):
        part = part.strip()
        if ":" in part:
            start, end = part.split(":", 1)
            rows.update(range(int(start), int(end) + 1))
        else:
            rows.add(int(part))
    return rows


def cell_display_value(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value)


def main():
    file_path = None
    sheet_name = None
    col_spec = None
    row_spec = None

    for arg in sys.argv[1:]:
        if arg.startswith("--sheet="):
            sheet_name = arg[len("--sheet="):]
        elif arg.startswith("--cols="):
            col_spec = arg[len("--cols="):]
        elif arg.startswith("--rows="):
            row_spec = arg[len("--rows="):]
        elif not arg.startswith("-"):
            file_path = arg
        else:
            print(f"Warning: unknown argument '{arg}'", file=sys.stderr)

    if not file_path:
        print(
            "Usage: python dump-xlsx.py <file> [--sheet=<name>] [--cols=A:C] [--rows=1:10]",
            file=sys.stderr,
        )
        sys.exit(1)

    # read_only + data_only for performance and to get computed formula values
    wb = load_workbook(file_path, read_only=True, data_only=True)

    if sheet_name is None:
        for name in wb.sheetnames:
            ws = wb[name]
            if ws.max_row and ws.max_column:
                col_label = get_column_letter(ws.max_column)
                print(f"{name} ({col_label} x {ws.max_row})", file=sys.stderr)
            else:
                print(f"{name} (empty)", file=sys.stderr)
        wb.close()
        return

    if sheet_name not in wb.sheetnames:
        try:
            idx = int(sheet_name)
        except ValueError:
            idx = 0
        if 1 <= idx <= len(wb.sheetnames):
            sheet_name = wb.sheetnames[idx - 1]
        else:
            print(f"Error: Sheet '{sheet_name}' not found.", file=sys.stderr)
            print(f"Available sheets: {', '.join(wb.sheetnames)}", file=sys.stderr)
            wb.close()
            sys.exit(1)

    ws = wb[sheet_name]
    col_indices = parse_col_ranges(col_spec) if col_spec else None
    row_numbers = parse_row_ranges(row_spec) if row_spec else None
    max_row = max(row_numbers) if row_numbers else None

    writer = csv.writer(sys.stdout)

    for row_num, row in enumerate(ws.iter_rows(), start=1):
        if max_row and row_num > max_row:
            break
        if row_numbers and row_num not in row_numbers:
            continue

        if col_indices:
            values = [
                cell_display_value(cell)
                for col_idx, cell in enumerate(row, start=1)
                if col_idx in col_indices
            ]
        else:
            values = [cell_display_value(cell) for cell in row]

        writer.writerow(values)

    wb.close()


if __name__ == "__main__":
    main()
